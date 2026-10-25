from flask import Flask, render_template, jsonify, send_from_directory
from pathlib import Path
import mimetypes
from datetime import datetime, date, time
from openpyxl import load_workbook
import unicodedata
from urllib.parse import quote
import configparser
import shutil
import tempfile
import threading
import os

# ------------------------------------------------------------
# CORREÇÃO TV: garante o MIME correto dos vídeos.
# No Windows o registro às vezes devolve tipo errado para .mp4,
# e navegadores de TV (Tizen/webOS/Android TV) se recusam a
# tocar um vídeo que não chega como "video/mp4".
# ------------------------------------------------------------
mimetypes.add_type("video/mp4", ".mp4")
mimetypes.add_type("video/mp4", ".m4v")
mimetypes.add_type("video/webm", ".webm")
mimetypes.add_type("video/ogg", ".ogv")
mimetypes.add_type("video/quicktime", ".mov")
mimetypes.add_type("image/webp", ".webp")
mimetypes.add_type("image/svg+xml", ".svg")

app = Flask(__name__)

BASE_DIR = Path(__file__).resolve().parent

# ------------------------------------------------------------
# CAMINHOS (configuráveis no arquivo config.ini)
# A planilha e a pasta de mídia podem ficar numa pasta
# compartilhada da rede, SEPARADAS do código. Assim os usuários
# editam avisos/salas sem ter acesso ao app.py nem ao HTML.
# ------------------------------------------------------------
CONFIG_FILE = BASE_DIR / "config.ini"


def ler_caminho(chave, padrao):
    config = configparser.ConfigParser()
    try:
        config.read(CONFIG_FILE, encoding="utf-8-sig")
        valor = config.get("caminhos", chave, fallback="").strip().strip('"')
    except Exception as erro:
        print("⚠ Erro ao ler config.ini:", erro)
        valor = ""
    if not valor:
        return padrao
    caminho = Path(os.path.expandvars(valor))
    if not caminho.is_absolute():
        caminho = BASE_DIR / caminho
    return caminho


# Planilha na unidade de rede Z: (pode ser trocada no config.ini)
PLANILHA_REDE = Path(
    r"Z:\PARTICULAR\Wallace Miguel\Painel de Avisos\avisos.xlsx"
)

def ler_lista_planilhas():
    """
    A linha "planilha" do config.ini aceita mais de um caminho,
    separados por  |   O painel usa o primeiro que existir.
    Ex.: Z:\\...\\avisos.xlsx | \\\\SERVIDOR\\Rede\\...\\avisos.xlsx
    """
    config = configparser.ConfigParser()
    try:
        config.read(CONFIG_FILE, encoding="utf-8-sig")
        bruto = config.get("caminhos", "planilha", fallback="")
    except Exception:
        bruto = ""
    caminhos = []
    for parte in bruto.split("|"):
        parte = parte.strip().strip('"')
        if not parte:
            continue
        c = Path(os.path.expandvars(parte))
        if not c.is_absolute():
            c = BASE_DIR / c
        caminhos.append(c)
    return caminhos or [PLANILHA_REDE]


PLANILHAS = ler_lista_planilhas()
EXCEL_FILE = PLANILHAS[0]


def escolher_planilha():
    """Aponta EXCEL_FILE para o primeiro caminho acessível."""
    global EXCEL_FILE
    for caminho in PLANILHAS:
        try:
            if caminho.exists():
                if caminho != EXCEL_FILE:
                    print("📄 Usando a planilha:", caminho)
                EXCEL_FILE = caminho
                return True
        except Exception:
            continue
    return False

# Cópia de segurança local: atualizada a cada leitura bem-sucedida da
# planilha da rede. Se o PC ligar e a rede ainda não estiver disponível,
# o painel abre com esta cópia em vez de ficar vazio.
BACKUP_FILE = BASE_DIR / "dados" / "avisos_backup.xlsx"
MEDIA_DIR = ler_caminho("midia", BASE_DIR / "static" / "avisos")

TRAVA_CACHE = threading.Lock()

VIDEO_EXTENSIONS = {".mp4", ".m4v", ".webm", ".ogg", ".ogv", ".mov"}

IMAGE_EXTENSIONS = {
    ".jpg", ".jpeg", ".png", ".gif",
    ".webp", ".bmp", ".svg"
}

CACHE = {
    "mtime": None,
    "avisos": [],
    "salas": [],
    "configuracoes": {}
}


# ============================================================
# FUNÇÕES AUXILIARES
# ============================================================

def texto(valor):
    if valor is None:
        return ""
    return str(valor).strip()


def normalizar(texto_valor):
    texto_valor = texto(texto_valor)
    texto_valor = unicodedata.normalize("NFD", texto_valor)
    texto_valor = "".join(
        c for c in texto_valor
        if unicodedata.category(c) != "Mn"
    )
    return "".join(
        c.lower() for c in texto_valor
        if c.isalnum()
    )


def valor_bool(valor):
    if valor is None:
        return True
    if texto(valor) == "":
        return True
    valor_normalizado = normalizar(valor)
    return valor_normalizado in {
        "sim", "s", "1", "true", "verdadeiro",
        "ativo", "yes", "y"
    }


def serializar(valor):
    if isinstance(valor, datetime):
        return valor.isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, time):
        return valor.strftime("%H:%M:%S")
    return valor


def detectar_tipo_midia(nome):
    extensao = Path(texto(nome)).suffix.lower()
    if extensao in VIDEO_EXTENSIONS:
        return "video"
    if extensao in IMAGE_EXTENSIONS:
        return "imagem"
    return "imagem"


def arquivo_existe(nome):
    if not nome:
        return False
    caminho = MEDIA_DIR / str(nome)
    return caminho.exists()


def verificar_video_para_tv(nome):
    """
    Verifica se o MP4 está com o índice (átomo 'moov') no início
    do arquivo ("faststart"). Sem isso, o navegador da TV precisa
    baixar o arquivo inteiro antes de começar a tocar e, na prática,
    o vídeo fica preto / não passa.
    Retorna uma mensagem de alerta ou "" se estiver tudo certo.
    """
    caminho = MEDIA_DIR / str(nome)
    if caminho.suffix.lower() not in {".mp4", ".m4v", ".mov"}:
        return ""
    if not caminho.exists():
        return ""
    try:
        with open(caminho, "rb") as f:
            posicao = 0
            tamanho_total = caminho.stat().st_size
            while posicao < tamanho_total:
                f.seek(posicao)
                cabecalho = f.read(16)
                if len(cabecalho) < 8:
                    break
                tamanho = int.from_bytes(cabecalho[0:4], "big")
                tipo = cabecalho[4:8]
                if tamanho == 1 and len(cabecalho) >= 16:
                    tamanho = int.from_bytes(cabecalho[8:16], "big")
                elif tamanho == 0:
                    tamanho = tamanho_total - posicao
                if tipo == b"moov":
                    return ""
                if tipo == b"mdat":
                    return (
                        "índice (moov) no FINAL do arquivo - a TV não "
                        "consegue tocar. Rode converter_videos.bat"
                    )
                if tamanho < 8:
                    break
                posicao += tamanho
    except Exception:
        return ""
    return ""


# ============================================================
# DATA
# ============================================================

def converter_data(valor):
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    valor = texto(valor)
    if not valor:
        return None
    formatos = [
        "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d",
        "%d/%m/%y", "%d-%m-%y"
    ]
    for formato in formatos:
        try:
            return datetime.strptime(valor, formato).date()
        except ValueError:
            continue
    return None


def formatar_data_hora_sala(valor):
    """
    Recebe o valor bruto da célula Data e devolve
    uma string amigável: dd/mm/aaaa
    """
    if valor is None or texto(valor) == "":
        return ""

    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")

    data = converter_data(valor)
    if data:
        return data.strftime("%d/%m/%Y")

    return texto(valor)


def formatar_hora_sala(valor):
    """
    Recebe o valor bruto da célula de Hora e devolve HH:MM
    """
    if valor is None or texto(valor) == "":
        return ""

    if isinstance(valor, time):
        return valor.strftime("%H:%M")
    if isinstance(valor, datetime):
        return valor.strftime("%H:%M")

    h = converter_hora(valor)
    if h:
        return h.strftime("%H:%M")

    return texto(valor)


# ============================================================
# HORA
# ============================================================

def converter_hora(valor):
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.time()
    if isinstance(valor, time):
        return valor
    valor = texto(valor)
    if not valor:
        return None
    formatos = ["%H:%M", "%H:%M:%S"]
    for formato in formatos:
        try:
            return datetime.strptime(valor, formato).time()
        except ValueError:
            continue
    return None


# ============================================================
# VERIFICAÇÃO DE PERÍODO
# ============================================================

def dentro_do_periodo(
    data_inicial=None,
    hora_inicial=None,
    data_final=None,
    hora_final=None
):
    agora = datetime.now()

    if data_inicial:
        inicio = datetime.combine(
            data_inicial,
            hora_inicial or time(0, 0, 0)
        )
        if agora < inicio:
            return False

    if data_final:
        fim = datetime.combine(
            data_final,
            hora_final or time(23, 59, 59)
        )
        if agora > fim:
            return False

    return True


# ============================================================
# LOCALIZAR COLUNA
# ============================================================

def encontrar_coluna(colunas, nomes):
    nomes_normalizados = {normalizar(nome) for nome in nomes}
    for indice, original in colunas.items():
        if normalizar(original) in nomes_normalizados:
            return indice
    return None


# ============================================================
# FORMATAR DATA
# ============================================================

def formatar_data(valor):
    data = converter_data(valor)
    if not data:
        return ""
    return data.strftime("%d/%m/%Y")


# ============================================================
# AVISOS
# ============================================================

def carregar_avisos(wb):

    if "Avisos" in wb.sheetnames:
        ws = wb["Avisos"]
    else:
        print("Aba 'Avisos' não encontrada.")
        return []

    linhas = list(ws.iter_rows(values_only=True))
    if not linhas:
        return []

    cabecalhos = linhas[0]
    colunas = {}
    for indice, nome in enumerate(cabecalhos):
        if nome is not None:
            colunas[indice] = texto(nome)

    col_id = encontrar_coluna(colunas, ["ID", "Código", "Codigo"])
    col_ativo = encontrar_coluna(colunas, ["Ativo", "Ativa", "Exibir", "Publicado"])
    col_titulo = encontrar_coluna(colunas, ["Título", "Titulo", "Assunto", "Nome"])
    col_descricao = encontrar_coluna(colunas, ["Descrição", "Descricao", "Texto", "Mensagem", "Conteúdo", "Conteudo"])
    col_imagem = encontrar_coluna(colunas, ["Imagem", "Arquivo", "Mídia", "Midia", "Vídeo", "Video", "Arquivo de Mídia", "Arquivo de Midia"])
    col_ordem = encontrar_coluna(colunas, ["Ordem", "Posição", "Posicao"])

    # ========================================================
    # COLUNA POSIÇÃO (novo)
    # ========================================================
    col_posicao = encontrar_coluna(
        colunas,
        [
            "Posicao", "Posição",
            "Local", "Localizacao", "Localização",
            "Destino", "Onde Exibir",
            "Categoria", "Tipo", "Aba", "Destaque",
            "Posição no Painel", "Posicao no Painel"
        ]
    )

    col_status = encontrar_coluna(colunas, ["Status", "Situação", "Situacao", "Estado"])
    col_responsavel = encontrar_coluna(colunas, ["Responsável", "Responsavel", "Autor", "Criado por"])
    col_data = encontrar_coluna(colunas, ["Data", "Data do Aviso", "Data Aviso", "Data de Publicação", "Data Publicacao", "Data Publicação"])
    col_data_inicial = encontrar_coluna(colunas, ["Data Inicial", "Data Inicial do Aviso", "Inicio", "Início"])
    col_hora_inicial = encontrar_coluna(colunas, ["Hora Inicial", "Hora Inicio", "Hora Início"])
    col_data_final = encontrar_coluna(colunas, ["Data Final", "Data Final do Aviso", "Fim"])
    col_hora_final = encontrar_coluna(colunas, ["Hora Final", "Hora Fim"])

    # Debug: mostra as colunas detectadas
    print(f"🔎 Coluna Posicao detectada: {col_posicao}")

    def pegar(linha, coluna):
        if coluna is None:
            return ""
        if coluna >= len(linha):
            return ""
        return linha[coluna]

    def normalizar_posicao(valor):
        v = normalizar(valor)
        if not v:
            return "carrossel"
        if v in {"carrossel", "carrosel", "carrousel", "carousel",
                 "slide", "slides"}:
            return "carrossel"
        if v in {"card1", "aviso1", "inferior1", "baixo1",
                 "card01", "card", "destaque"}:
            return "card1"
        if v in {"card2", "aviso2", "inferior2", "baixo2", "card02"}:
            return "card2"
        return "carrossel"

    avisos = []

    for linha_numero, linha in enumerate(linhas[1:], start=2):
        ativo_valor = pegar(linha, col_ativo)
        if col_ativo is not None:
            if not valor_bool(ativo_valor):
                continue

        titulo = texto(pegar(linha, col_titulo))
        descricao = texto(pegar(linha, col_descricao))
        arquivo = texto(pegar(linha, col_imagem))

        if not arquivo:
            continue

        data_inicial = converter_data(pegar(linha, col_data_inicial))
        hora_inicial = converter_hora(pegar(linha, col_hora_inicial))
        data_final = converter_data(pegar(linha, col_data_final))
        hora_final = converter_hora(pegar(linha, col_hora_final))

        if not dentro_do_periodo(data_inicial, hora_inicial, data_final, hora_final):
            continue

        status = texto(pegar(linha, col_status)) or "Publicado"
        responsavel = texto(pegar(linha, col_responsavel))

        data_aviso = pegar(linha, col_data)
        if data_aviso:
            data_formatada = formatar_data(data_aviso)
        elif data_inicial:
            data_formatada = data_inicial.strftime("%d/%m/%Y")
        else:
            data_formatada = ""

        ordem_valor = pegar(linha, col_ordem)
        try:
            ordem = int(ordem_valor)
        except (ValueError, TypeError):
            ordem = 9999

        tipo = detectar_tipo_midia(arquivo)
        existe = arquivo_existe(arquivo)

        if not existe:
            print(f"⚠ Arquivo não encontrado em static/avisos: {arquivo}")
        elif tipo == "video":
            alerta = verificar_video_para_tv(arquivo)
            if alerta:
                print(f"⚠ Vídeo {arquivo}: {alerta}")
        media_url = "/avisos/" + quote(arquivo, safe="/")

        posicao_bruta = pegar(linha, col_posicao)
        posicao = normalizar_posicao(posicao_bruta)

        avisos.append({
            "id": texto(pegar(linha, col_id)) or str(linha_numero),
            "ativo": True,
            "titulo": titulo,
            "descricao": descricao,
            "imagem": arquivo,
            "arquivo": arquivo,
            "tipo": tipo,
            "media_tipo": tipo,
            "media_url": media_url,
            "ordem": ordem,
            "arquivo_existe": existe,
            "status": status,
            "responsavel": responsavel,
            "data": data_formatada,
            "data_aviso": data_formatada,
            "posicao": posicao,
            "data_inicial": data_inicial.isoformat() if data_inicial else "",
            "hora_inicial": hora_inicial.strftime("%H:%M") if hora_inicial else "",
            "data_final": data_final.isoformat() if data_final else "",
            "hora_final": hora_final.strftime("%H:%M") if hora_final else ""
        })

    avisos.sort(key=lambda x: (x["ordem"], x["id"]))
    print(f"✔ {len(avisos)} aviso(s) carregado(s)")
    return avisos

# ============================================================
# SALAS
# ============================================================

def carregar_salas(wb):

    if "Salas" not in wb.sheetnames:
        print("Aba 'Salas' não encontrada.")
        return []

    ws = wb["Salas"]
    linhas = list(ws.iter_rows(values_only=True))

    if not linhas:
        return []

    cabecalhos = [normalizar(x) for x in linhas[0]]

    print("📋 Cabeçalhos da aba Salas:", cabecalhos)

    salas = []

    for linha in linhas[1:]:

        item = {}
        item_raw = {}

        for i, valor in enumerate(linha):
            if i < len(cabecalhos):
                chave = cabecalhos[i]
                item[chave] = serializar(valor)
                item_raw[chave] = valor

        # ---------------------------------------------------
        # NOME DA SALA
        # ---------------------------------------------------
        nome_sala = texto(item.get("sala", ""))

        if not nome_sala:
            continue

        if nome_sala.strip() in {"0", "-", "—", "N/A"}:
            continue

        item["nome"] = nome_sala
        item["sala"] = nome_sala

        # ---------------------------------------------------
        # DATA FORMATADA
        # ---------------------------------------------------
        item["data_formatada"] = formatar_data_hora_sala(
            item_raw.get("data")
        )

        # ---------------------------------------------------
        # HORAS FORMATADAS
        # ---------------------------------------------------
        item["hora_inicial_fmt"] = formatar_hora_sala(
            item_raw.get("horainicial")
        )
        item["hora_final_fmt"] = formatar_hora_sala(
            item_raw.get("horafinal")
        )

        # ---------------------------------------------------
        # STATUS NORMALIZADO
        # ---------------------------------------------------
        status_original = texto(item.get("status", "")).upper()

        if any(p in status_original for p in
               ["OCUP", "EM USO", "USANDO"]):
            item["status"] = "OCUPADA"
        elif any(p in status_original for p in
                 ["RESERV", "AGENDAD"]):
            item["status"] = "RESERVADA"
        elif any(p in status_original for p in
                 ["LIVRE", "DISPON", "LIBERAD", "VAZIA"]):
            item["status"] = "DISPONIVEL"
        elif not status_original:
            item["status"] = "DISPONIVEL"
        else:
            item["status"] = status_original

        salas.append(item)

    print(f"✔ {len(salas)} sala(s) carregada(s)")
    if salas:
        print("🔎 Exemplo de sala:", salas[0])

    return salas


# ============================================================
# CONFIGURAÇÕES
# ============================================================

def carregar_configuracoes(wb):

    if "Configurações" not in wb.sheetnames:
        return {}

    ws = wb["Configurações"]
    configuracoes = {}

    for linha in ws.iter_rows(values_only=True):
        if len(linha) >= 2:
            chave = texto(linha[0])
            valor = serializar(linha[1])
            if chave:
                configuracoes[chave] = valor

    return configuracoes


# ============================================================
# CACHE
# ============================================================

def obter_mtime():
    if not EXCEL_FILE.exists():
        return None
    try:
        stat = EXCEL_FILE.stat()
        return (stat.st_mtime_ns, stat.st_size)
    except Exception:
        return None


def abrir_planilha(origem=None):
    """
    Copia a planilha para uma pasta temporária e abre a CÓPIA.
    - Não trava o arquivo da rede (quem está editando consegue salvar).
    - Se o Excel estiver no meio do salvamento, a leitura falha e
      o painel mantém os dados anteriores; tenta de novo em seguida.
    """
    pasta_temp = tempfile.mkdtemp(prefix="painel_")
    copia = Path(pasta_temp) / "avisos_copia.xlsx"
    try:
        shutil.copy2(origem or EXCEL_FILE, copia)
        wb = load_workbook(copia, data_only=True, read_only=False)
        if origem is None:
            # leitura da rede deu certo -> atualiza a cópia de segurança
            try:
                BACKUP_FILE.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(copia, BACKUP_FILE)
            except Exception as erro:
                print("⚠ Não foi possível salvar a cópia de segurança:", erro)
        return wb
    finally:
        shutil.rmtree(pasta_temp, ignore_errors=True)


def carregar_backup_local():
    """Usada só quando a rede está fora e o painel ainda está vazio."""
    if CACHE["mtime"] is not None or CACHE.get("backup_carregado"):
        return
    if not BACKUP_FILE.exists():
        return
    try:
        wb = abrir_planilha(BACKUP_FILE)
        CACHE["avisos"] = carregar_avisos(wb)
        CACHE["salas"] = carregar_salas(wb)
        CACHE["configuracoes"] = carregar_configuracoes(wb)
        wb.close()
        CACHE["backup_carregado"] = True
        print("💾 Rede indisponível - exibindo a cópia de segurança local.")
    except Exception as erro:
        print("⚠ Não foi possível ler a cópia de segurança:", erro)


def atualizar_cache():
    with TRAVA_CACHE:
        existe = escolher_planilha()

        if not existe:
            if not CACHE.get("aviso_rede"):
                print("⚠ Planilha não encontrada em nenhum destes caminhos:")
                for c in PLANILHAS:
                    print("     ", c)
                print("   Dicas: 1) abra o caminho no Explorer deste PC;")
                print("          2) NÃO rode o painel 'como administrador'")
                print("             (o Windows esconde a unidade Z: nesse modo);")
                print("          3) rode  net use  para ver o caminho real do Z:")
                CACHE["aviso_rede"] = True
            carregar_backup_local()
            return

        if CACHE.get("aviso_rede"):
            print("✔ Planilha da rede acessível novamente.")
            CACHE["aviso_rede"] = False

        mtime = obter_mtime()
        if mtime is None or CACHE["mtime"] == mtime:
            return

        print("🔄 Detectada alteração no Excel. Recarregando...")
        try:
            wb = abrir_planilha()
            avisos_novos = carregar_avisos(wb)
            salas_novas = carregar_salas(wb)
            config_novas = carregar_configuracoes(wb)
            wb.close()
        except Exception as erro:
            # arquivo sendo salvo / rede instável: NÃO apaga o painel,
            # mantém o que estava na tela e tenta na próxima consulta
            print("⚠ Não foi possível ler a planilha agora:", erro)
            return

        CACHE["avisos"] = avisos_novos
        CACHE["salas"] = salas_novas
        CACHE["configuracoes"] = config_novas
        CACHE["mtime"] = mtime
        print(
            f"[cache] avisos={len(CACHE['avisos'])} | "
            f"salas={len(CACHE['salas'])}"
        )


# ============================================================
# ROTAS
# ============================================================

@app.route("/")
def index():
    atualizar_cache()
    return render_template("index.html")


@app.route("/api/avisos")
def api_avisos():
    atualizar_cache()
    return jsonify(CACHE["avisos"])


@app.route("/api/salas")
def api_salas():
    atualizar_cache()
    return jsonify(CACHE["salas"])


@app.route("/api/configuracoes")
def api_configuracoes():
    atualizar_cache()
    return jsonify(CACHE["configuracoes"])


@app.route("/api/painel")
def api_painel():
    atualizar_cache()
    return jsonify({
        "avisos": CACHE["avisos"],
        "salas": CACHE["salas"],
        "configuracoes": CACHE["configuracoes"]
    })


# ============================================================
# ATENDIMENTO (Messenger) - lido do arquivo gerado pelo
# leitor_messenger.py
# ============================================================

ATENDIMENTO_JSON = BASE_DIR / "dados" / "atendimento.json"
ATENDIMENTO_VALIDADE = 90   # segundos sem atualizar = leitor parado


@app.route("/api/atendimento")
def api_atendimento():
    import json
    import time as _time

    if not ATENDIMENTO_JSON.exists():
        return jsonify({"status": "sem_leitor",
                        "mensagem": "O leitor do Messenger ainda não rodou neste PC"})
    try:
        with open(ATENDIMENTO_JSON, "r", encoding="utf-8") as f:
            dados = json.load(f)
    except Exception:
        return jsonify({"status": "carregando"})

    if dados.get("status") == "desativado":
        return jsonify(dados)

    idade = _time.time() - float(dados.get("atualizado_ts") or 0)
    if idade > ATENDIMENTO_VALIDADE:
        dados["status"] = "sem_leitor"
        dados["mensagem"] = ("O leitor do Messenger parou "
                             "(janela 'Leitor Messenger' fechada ou com erro)")
    return jsonify(dados)


# ============================================================
# DIAGNÓSTICO  ->  abra  http://IP-DO-PC:5000/diagnostico
# ============================================================

@app.route("/diagnostico")
def diagnostico():
    import html
    import json
    linhas = []

    linhas.append("=== PLANILHA ===")
    for c in PLANILHAS:
        try:
            ok = c.exists()
        except Exception:
            ok = False
        linhas.append(f"{'OK ' if ok else 'NÃO'}  {c}")
    linhas.append(f"Em uso: {EXCEL_FILE}")
    linhas.append(f"Cópia de segurança: {BACKUP_FILE} "
                  f"({'existe' if BACKUP_FILE.exists() else 'não existe'})")

    linhas.append("")
    linhas.append("=== MESSENGER (dados/atendimento.json) ===")
    if ATENDIMENTO_JSON.exists():
        try:
            dados = json.loads(ATENDIMENTO_JSON.read_text(encoding="utf-8"))
            import time as _t
            idade = int(_t.time() - float(dados.get("atualizado_ts") or 0))
            linhas.append(f"Atualizado há {idade} s")
            linhas.append(json.dumps(dados, ensure_ascii=False, indent=2))
        except Exception as erro:
            linhas.append(f"Erro lendo: {erro}")
    else:
        linhas.append("Arquivo não existe -> o leitor nunca rodou neste PC.")

    log = BASE_DIR / "dados" / "leitor.log"
    linhas.append("")
    linhas.append("=== ÚLTIMAS LINHAS DO LEITOR (dados/leitor.log) ===")
    if log.exists():
        try:
            texto_log = log.read_text(encoding="utf-8", errors="ignore").splitlines()
            linhas.extend(texto_log[-40:])
        except Exception as erro:
            linhas.append(f"Erro lendo: {erro}")
    else:
        linhas.append("Sem log -> o leitor não foi iniciado.")

    corpo = html.escape("\n".join(linhas))
    return (
        "<!doctype html><meta charset='utf-8'><title>Diagnóstico</title>"
        "<body style='background:#111;color:#ddd;font:14px monospace;padding:20px'>"
        "<h2 style='color:#f0a04b'>Diagnóstico do Painel</h2>"
        f"<pre style='white-space:pre-wrap'>{corpo}</pre></body>"
    )


@app.route("/avisos/<path:nome_arquivo>")
def avisos(nome_arquivo):
    # conditional=True -> responde a pedidos "Range" (HTTP 206).
    # Navegadores de TV SÓ tocam vídeo se o servidor aceitar Range.
    # max_age -> a TV guarda o vídeo em cache e não baixa tudo de
    # novo a cada volta do carrossel.
    resposta = send_from_directory(
        MEDIA_DIR,
        nome_arquivo,
        conditional=True,
        max_age=3600
    )
    resposta.headers["Accept-Ranges"] = "bytes"
    return resposta


# ============================================================
# EXECUÇÃO
# ============================================================

if __name__ == "__main__":
    try:
        MEDIA_DIR.mkdir(parents=True, exist_ok=True)
    except Exception as erro:
        print("⚠ Não foi possível acessar a pasta de mídia:", erro)

    print("=" * 55)
    print("⚙ Configuração:", CONFIG_FILE,
          "(encontrado)" if CONFIG_FILE.exists() else "(não existe, usando pastas locais)")
    print("📂 Excel esperado em:", EXCEL_FILE)
    print("📂 Existe?", EXCEL_FILE.exists())
    print("📂 Mídia em:", MEDIA_DIR)
    print("=" * 55)

    # threaded=True  -> atende vários pedidos ao mesmo tempo
    #                   (a TV abre várias conexões para o vídeo).
    # debug=False     -> o modo debug reinicia o servidor sozinho e
    #                   corta o vídeo no meio; além disso não é seguro
    #                   deixar ligado na rede.
    app.run(
        host="0.0.0.0",
        port=5000,
        debug=False,
        threaded=True
    )